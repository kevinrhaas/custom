#!/usr/bin/env python3
"""deterministic_xlsx.py — make a written .xlsx a function of its DATA, not of the clock.

THE PROBLEM, MEASURED. `openpyxl` writes a workbook as a zip, and a zip stamps every
member with the time it was written. So two builds of the SAME data, seconds apart,
produce different bytes:

    build 1  b2a914407e8a7ec96156fda530d962c0
    build 2  d4805cc5c55eeca85ef84a97b3beda16

The .csv sibling beside it, from the same rows in the same run, is byte-identical both
times. The difference is not the data. It is the timestamps.

WHAT THAT COST. `.xlsx` is binary to git, so a file that changes on every build conflicts
on every merge and cannot be merged — only chosen. Two of them,
`resident_audit_master.xlsx` and `T-0509_resident_research_working.xlsx`, turned up in
conflict after conflict through 2026-09-17/18, and at least once the ENTIRE diff was
`docProps/core.xml`: same byte count, not one sheet different, nothing but the clock. The
owner, looking at yet another pair of them in a conflict list: "these spreadsheets
constantly cause a problem with needing resolution do we need to keep them around".

WHY NOT JUST DELETE THEM. Because the data is not the problem and it is worth having: each
workbook sits beside a .csv of the same rows and a README, and the .csv is what tools read.
The workbook is the human-readable copy. Making it reproducible keeps the deliverable and
removes the churn, which is strictly better than dropping it — and unlike a .gitattributes
merge driver it fixes the CAUSE, so an unchanged workbook stops appearing in `git status`
at all.

HOW. Rewrite the archive with every member's timestamp pinned and the volatile bits of
`docProps/core.xml` (created/modified) normalised, members in sorted order and deflate
level fixed. The cell data is untouched; Excel reads the result exactly as before.

    from deterministic_xlsx import settle
    book.save(path)
    settle(path)
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

# 1980-01-01 00:00:00 — the earliest a zip can express, so it reads as "no time" rather
# than as a date somebody might mistake for provenance.
EPOCH = (1980, 1, 1, 0, 0, 0)

_STAMP = re.compile(
    rb"<(dcterms:(?:created|modified)[^>]*)>[^<]*</(dcterms:(?:created|modified))>")


def _settle_core_props(raw: bytes) -> bytes:
    """Pin the created/modified stamps openpyxl writes from the clock."""
    return _STAMP.sub(rb"<\1>1980-01-01T00:00:00Z</\2>", raw)


def settle(path: str | Path) -> Path:
    """Rewrite an .xlsx in place so identical data gives identical bytes."""
    path = Path(path)
    with zipfile.ZipFile(path) as zf:
        members = sorted(zf.namelist())
        payload = {name: zf.read(name) for name in members}

    for name in list(payload):
        if name == "docProps/core.xml":
            payload[name] = _settle_core_props(payload[name])

    tmp = path.with_suffix(path.suffix + ".settling")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as out:
        for name in members:
            info = zipfile.ZipInfo(name, date_time=EPOCH)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            out.writestr(info, payload[name])
    tmp.replace(path)
    return path


def self_test() -> int:
    """Two workbooks with the same rows, written apart in time, must be byte-identical."""
    import hashlib, tempfile, time
    from openpyxl import Workbook

    def write(p):
        b = Workbook()
        s = b.active
        s.append(["name", "count"])
        s.append(["Elston", 2])
        b.save(p)
        return p

    fails = 0
    with tempfile.TemporaryDirectory() as d:
        a, b = Path(d) / "a.xlsx", Path(d) / "b.xlsx"
        write(a); time.sleep(2); write(b)
        raw_same = hashlib.md5(a.read_bytes()).digest() == hashlib.md5(b.read_bytes()).digest()
        print(f"   self-test | {'ok  ' if not raw_same else 'FAIL'} openpyxl alone is NOT "
              f"deterministic (this is the fault being fixed)")
        fails += 1 if raw_same else 0

        settle(a); settle(b)
        settled_same = a.read_bytes() == b.read_bytes()
        print(f"   self-test | {'ok  ' if settled_same else 'FAIL'} settled, the same rows give "
              f"the same bytes")
        fails += 0 if settled_same else 1

        settle(a)
        print(f"   self-test | {'ok  ' if a.read_bytes() == b.read_bytes() else 'FAIL'} "
              f"settling twice is a no-op")
        fails += 0 if a.read_bytes() == b.read_bytes() else 1

        from openpyxl import load_workbook
        rows = list(load_workbook(a).active.values)
        ok = rows == [("name", "count"), ("Elston", 2)]
        print(f"   self-test | {'ok  ' if ok else 'FAIL'} the cells still read back unchanged "
              f"— {rows}")
        fails += 0 if ok else 1
    print(f"   self-test | {fails} failure(s)")
    return fails


if __name__ == "__main__":
    import sys
    raise SystemExit(self_test() if "--self-test" in sys.argv else 0)
